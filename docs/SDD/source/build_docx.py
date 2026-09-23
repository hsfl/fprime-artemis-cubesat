"""Build the editable Word document from Artemis_SDD.md and diagram previews.
Requires python-docx. Run with the bundled Python documented in ../README.md.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT/'source/Artemis_SDD.md'
OUT=ROOT/'Artemis_CubeSat_Kit_SDD_v1.0.docx'
sections=SRC.read_text().split('---PAGE---')
doc=Document()
doc.settings.odd_and_even_pages_header_footer=False
for st in doc.styles:
 for el in list(st.element.iter(qn('w:pBdr'))): el.getparent().remove(el)
sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65)
sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
sec.header_distance=Inches(.28);sec.footer_distance=Inches(.28)
for name in ['Normal','Body Text','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
 s=doc.styles[name];s.font.name='Arial';s.font.color.rgb=RGBColor(0,0,0)
 s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial')
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
doc.styles['Normal'].paragraph_format.left_indent=Inches(0)
doc.styles['Normal'].paragraph_format.first_line_indent=Inches(0)
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
for n,size in [('Title',30),('Subtitle',19),('Heading 1',19),('Heading 2',12),('Heading 3',11)]:
 s=doc.styles[n];s.font.size=Pt(size);s.font.bold=True
 s.paragraph_format.space_before=Pt(10);s.paragraph_format.space_after=Pt(8)
 s.paragraph_format.keep_with_next=True
s=doc.styles['Caption'];s.font.size=Pt(9);s.font.italic=True;s.paragraph_format.space_after=Pt(9)
head=sec.header.paragraphs[0];head.text='ARTEMIS CUBESAT KIT  |  SYSTEM AND SOFTWARE DESIGN';head.style='Normal'
for r in head.runs:r.font.size=Pt(8)
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
foot.add_run('Version 1.0  •  Design baseline     ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
for r in foot.runs:r.font.size=Pt(8)
doc.core_properties.title='Artemis CubeSat Kit System and Software Design Document'
doc.core_properties.subject='Reusable bus architecture and mission payload extension contracts'
doc.core_properties.author='Artemis CubeSat Kit'
doc.core_properties.keywords='F Prime, Zephyr, Teensy, Raspberry Pi, Yamcs, SDD'
bookmarkid=0
def bookmark(p,name):
 global bookmarkid
 bookmarkid+=1
 a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(bookmarkid));a.set(qn('w:name'),name)
 b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),str(bookmarkid))
 p._p.insert(0,a);p._p.append(b)
def link(p,label,target,internal=False):
 h=OxmlElement('w:hyperlink')
 if internal:h.set(qn('w:anchor'),target)
 else:h.set(qn('r:id'),p.part.relate_to(target,RT.HYPERLINK,is_external=True))
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'1C4C73');pr.append(c);r.append(pr)
 t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)
def rich(p,text):
 for token in re.split(r'(\*\*.*?\*\*|`[^`]+`|https?://\S+)',text):
  if not token:continue
  if token.startswith('**'):p.add_run(token[2:-2]).bold=True
  elif token.startswith('`'):
   r=p.add_run(token[1:-1]);r.font.name='Arial';r.font.size=Pt(9)
  elif token.startswith('http'):link(p,token.rstrip(' .'),token.rstrip(' .'))
  else:p.add_run(token)
def table(rows):
 t=doc.add_table(rows=1,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 n=len(rows[0]);weights=([.25,.75] if n==2 else [.23,.36,.41])
 if rows[0][0]=='ID':weights=[.09,.40,.51]
 widths=[6.9*w for w in weights]
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for i,row in enumerate(rows):
  cells=t.rows[0].cells if i==0 else t.add_row().cells
  trpr=cells[0]._tc.getparent().get_or_add_trPr()
  cant=OxmlElement('w:cantSplit');trpr.append(cant)
  if i==0:
   repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
  for j,txt in enumerate(row):
   c=cells[j];c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   p=c.paragraphs[0];p.paragraph_format.space_after=Pt(1 if section.strip().startswith(('# 21 ', '# 22 ')) else 3);p.paragraph_format.space_before=Pt(1 if section.strip().startswith(('# 21 ', '# 22 ')) else 3);p.paragraph_format.line_spacing=1.03
   rich(p,txt)
   for r in p.runs:r.font.size=Pt(9);r.bold=(i==0)
   tcpr=c._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6EF' if i==0 else ('F3F5F7' if i%2==0 else 'FFFFFF'));tcpr.append(shade)
   mar=OxmlElement('w:tcMar')
   for edge,val in [('top','75'),('bottom','75'),('left','95'),('right','95')]:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'), '35' if edge in ('top','bottom') and section.strip().startswith(('# 21 ', '# 22 ')) else val);e.set(qn('w:type'),'dxa');mar.append(e)
   tcpr.append(mar)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   tcpr.append(borders)
 doc.add_paragraph().paragraph_format.space_after=Pt(1)
 return t
source_targets={
 'S1':'../../../artemis-cubesat-kit-official-textbook_A-Guide-to-CubeSat-Mission-and-Bus-Design-1696282662.pdf',
 'S2':'../../../fprime-material/fprime/docs/user-manual/index.md',
 'S5':'../../../artemis-cubesat-pdu-firmware/PDU_PROTOCOL_ICD.md',
 'S6':'../../../docs/context/README.md',
 'S8':'../../../fprime-material/fprime-yamcs/README.md',
 'S10':'../../../artemis-hardware/Eagle Designs/OBC/OBC_V4.24/OBCv4.24.sch'}
for idx,section in enumerate(sections):
 # Page breaks belong to headings so no empty break paragraph can spill.
 lines=section.strip().splitlines();i=0
 while i<len(lines):
  line=lines[i].strip();i+=1
  if not line:continue
  if line=='{{TOC}}':
   for k,part in enumerate(sections[2:],start=3):
    title=part.strip().splitlines()[0].removeprefix('# ')
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.0
    link(p,title,'sec'+str(k-2),True)
    p.add_run('  '+str(k)).font.size=Pt(10)
   continue
  if line.startswith('{{FIG:'):
   name,cap=line[6:-2].split('|',1)
   matches=list((ROOT/'diagrams').glob(name[:2]+'_*.png'))
   if not matches:raise RuntimeError('Missing figure '+name)
   p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
   r=p.add_run();shape=r.add_picture(str(matches[0]),width=Inches(6.9))
   shape._inline.docPr.set('descr',cap)
   doc.add_paragraph(cap,'Caption');continue
  if line.startswith('|'):
   rows=[]
   while True:
    row=[x.strip() for x in line.strip('|').split('|')]
    if not all(re.fullmatch(r'[-: ]+',x) for x in row):rows.append(row)
    if i>=len(lines) or not lines[i].strip().startswith('|'):break
    line=lines[i].strip();i+=1
   table(rows);continue
  if line.startswith('# '):
   p=doc.add_paragraph(line[2:],'Title' if idx==0 else 'Heading 1')
   if idx:p.paragraph_format.page_break_before=True
   if idx>=2:bookmark(p,'sec'+str(idx-1))
  elif line.startswith('## '):doc.add_paragraph(line[3:],'Subtitle' if idx==0 else 'Heading 2')
  elif line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2')
  elif line.startswith('- '):
   p=doc.add_paragraph(style='List Bullet');rich(p,line[2:]);p.paragraph_format.space_after=Pt(5)
  else:
   p=doc.add_paragraph(style='Normal');p.alignment=WD_ALIGN_PARAGRAPH.LEFT;rich(p,line)
   if section.strip().startswith('# 23 '):p.paragraph_format.space_after=Pt(2)
   m=re.match(r'\[(S\d+)\]',line)
   if m:
    bookmark(p,m[1])
    if m[1] in source_targets:p.add_run(' ');link(p,'Open local source',source_targets[m[1]])
# add document settings update fields
settings=doc.settings.element
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
doc.save(OUT)
print(OUT)
print('Designed pages:',len(sections))
